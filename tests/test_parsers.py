from __future__ import annotations

from pathlib import Path

import pytest

from hammerstein.parsers import FILE_READ_MAX_CHARS, extract_text


def _build_minimal_pdf(visible: str = "HelloPDF") -> bytes:
    """Build a tiny valid PDF with extractable text (no external fixture file)."""

    def esc(s: str) -> str:
        return s.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")

    stream_payload = f"BT /F1 12 Tf 10 700 Td ({esc(visible)}) Tj ET".encode("latin1")
    stream_body = (
        f"<< /Length {len(stream_payload)} >>\nstream\n".encode("ascii")
        + stream_payload
        + b"\nendstream"
    )
    bodies = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        stream_body,
    ]
    buf = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets: list[int] = [0] * (len(bodies) + 1)
    for i, body in enumerate(bodies, start=1):
        offsets[i] = len(buf)
        buf.extend(f"{i} 0 obj\n".encode("ascii"))
        buf.extend(body)
        buf.extend(b"\nendobj\n")
    xref_pos = len(buf)
    max_id = len(bodies)
    xref = bytearray()
    xref.extend(b"xref\n")
    xref.extend(f"0 {max_id + 1}\n".encode("ascii"))
    xref.extend(b"0000000000 65535 f \n")
    for i in range(1, max_id + 1):
        xref.extend(f"{offsets[i]:010d} 00000 n \n".encode("ascii"))
    buf.extend(xref)
    trailer = (
        f"trailer\n<< /Size {max_id + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_pos}\n%%EOF\n"
    )
    buf.extend(trailer.encode("ascii"))
    return bytes(buf)


def test_extract_md(tmp_path: Path) -> None:
    p = tmp_path / "note.md"
    p.write_text("hello", encoding="utf-8")
    assert extract_text(p) == "hello"


def test_extract_txt(tmp_path: Path) -> None:
    p = tmp_path / "note.txt"
    p.write_text("hello", encoding="utf-8")
    assert extract_text(p) == "hello"


def test_extract_csv(tmp_path: Path) -> None:
    p = tmp_path / "data.csv"
    p.write_text("name,age\nalice,30\nbob,25\n", encoding="utf-8")
    out = extract_text(p)
    assert "|" in out
    assert "alice" in out
    assert "30" in out
    assert "bob" in out
    assert "25" in out


def test_extract_pdf(tmp_path: Path) -> None:
    pytest.importorskip("pypdf")
    p = tmp_path / "sample.pdf"
    p.write_bytes(_build_minimal_pdf("HelloPDF"))
    out = extract_text(p)
    assert "HelloPDF" in out


def test_extract_docx(tmp_path: Path) -> None:
    pytest.importorskip("docx")
    from docx import Document

    p = tmp_path / "doc.docx"
    doc = Document()
    doc.add_paragraph("hello world")
    doc.save(p)
    out = extract_text(p)
    assert "hello world" in out


def test_extract_xlsx(tmp_path: Path) -> None:
    pytest.importorskip("openpyxl")
    from openpyxl import Workbook

    p = tmp_path / "wb.xlsx"
    wb = Workbook()
    ws_a = wb.active
    ws_a.title = "alpha"
    ws_a.append(["h1", "h2"])
    ws_a.append(["v1", "v2"])
    ws_b = wb.create_sheet("beta")
    ws_b.append(["x", "y"])
    ws_b.append(["a", "b"])
    wb.save(p)
    wb.close()

    out = extract_text(p)
    assert "## alpha" in out
    assert "## beta" in out
    assert "v1" in out
    assert "v2" in out
    assert "a" in out
    assert "b" in out


def test_utf8_sanitization(tmp_path: Path) -> None:
    """File with invalid UTF-8 bytes (e.g. Windows mojibake / lone surrogates
    from upstream encoding bugs) must round-trip through extract_text into a
    string that encodes cleanly as UTF-8 — this is the gate that prevents the
    wrapper logger crash that hit the ham-025 falsification test."""
    p = tmp_path / "bad.txt"
    p.write_bytes(b"hello \x80\x81 world")
    out = extract_text(p)
    out.encode("utf-8")
    assert "�" in out


def test_unsupported_extension() -> None:
    with pytest.raises(ValueError):
        extract_text(Path("foo.pages"))


def test_truncation(tmp_path: Path) -> None:
    p = tmp_path / "big.txt"
    p.write_text("a" * 20_000, encoding="utf-8")
    out = extract_text(p)
    assert "[...content truncated at 16K chars...]" in out
    assert len(out) == FILE_READ_MAX_CHARS + len("\n\n[...content truncated at 16K chars...]")
